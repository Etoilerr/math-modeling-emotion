# -*- coding: utf-8 -*-
"""整理网上坑点与思路分析为 Word 表格文档。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# 页面设置
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# 默认字体
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def set_cell(cell, text, bold=False, size=10.5, bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if bg:
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(qn("w:shd"), {qn("w:fill"): bg})
        shading.append(shd)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

# ============ 标题 ============
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("E 题网上坑点与思路分析整理")
run.font.name = "Times New Roman"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
run.font.size = Pt(18)
run.bold = True

doc.add_paragraph()

# ============ 第一部分：坑点表 ============
add_heading("一、14 个易踩坑点及本队应对", level=1)

pits = [
    ("1", "数据接口前后不一致",
     "附件2对齐版有 text(50×768) 和 text_bert(3×50)，但附件3对齐版只有 text_bert 没有 text；附件3未对齐版连 text_bert 都没有。训练验证和专项测试必须用同一接口。",
     "已修。全程用 text_bert 的 attention_mask 构造 mt，不再依赖 text 字段；附件3推理直接读 text_bert。"),
    ("2", "索引方式反直觉",
     "题目明确 data['train']['audio'][j]，而非 data['train'][j]['audio']。按后者写会 KeyError。",
     "正确。to_tensors 用 d['text']/d['audio']/d['vision'] 顶层键索引。"),
    ("3", "附件3无 id 字段",
     "附件3 pkl 内没有 'id' 键，只能靠文件名（附件3_01.pkl）对应样本。",
     "正确。predict_attachment3.py 用文件名 basename 作为 sample_id。"),
    ("4", "局部缺失≠整模态缺失",
     "题目反复强调局部缺失是部分连续时间段不可用，不是整个模态不存在。用整模态 dropout 训练会在附件3上答偏。",
     "已修。增强改为 audio+vision 在相同位置做连续短缺失（长度1-4），text 不做中间缺失。"),
    ("5", "附件3文本没被局部遮挡",
     "实测30条样本 text_bert 注意力掩码全是尾部填充型，没有中间断裂；局部全零只在 audio/vision，且29/30条 audio 与 vision 零掩码位置完全相同。",
     "已修。增强策略对齐实测：text 不做中间缺失，audio+vision 同时缺失。"),
    ("6", "训练集本身就有缺失",
     "附件2训练集 vision 存在中间零段（如某样本 vision 零区间为[0,5]、[8,10]、[12,49]）。全零段既可能是填充也可能是真缺失。",
     "正确。build_mask 按||x||>0自动识别有效位置，不假设全零段都是 padding。"),
    ("7", "位置0恒零，文本填充非零",
     "附件2 audio 第0个时间步100%为全零（固定现象）；text(768)填充位置不是零（实测填充段范数均值9.76），必须用 text_bert 注意力掩码做掩码池化。",
     "正确。masked_mean_pool 用掩码加权，附件3用 text_bert attention_mask。"),
    ("8", "0 只属于中性",
     "题目明确0仅属于中性，不能同时归入负向或正向。",
     "正确。reg<0→Negative, reg=0→Neutral, reg>0→Positive。"),
    ("9", "禁止外部情感数据集",
     "不得引入 MOSI/IEMOCAP/MELD 等外部情感数据集训练、微调、调参或统计。开源预训练模型可用，但不能用外部情感标注。",
     "正确。仅用附件2 CMU-MOSEI 数据，未引入外部情感标注。"),
    ("10", "附件3/4无标签，只做最终推理",
     "不能伪标签、不能用测试集调阈值。只能做最终推理。",
     "正确。附件3/4仅推理，阈值和超参只在验证集选。"),
    ("11", "可解释性必须落到原始时间轴",
     "题目要求证据对应原始文本片段、语音时段或视觉关键帧，但附件2不保存原始帧、波形和逐词时间戳。必须自建特征位置→原始时间的映射（附件4配套视频可用）。",
     "已做。用附件4视频抽关键帧，IG+遮挡+注意力三路互证。"),
    ("12", "交叉引用错误",
     "E题正文写'图3用一个独立样本示意'但全文只有一张图且题注是图1。照抄会编号错误。",
     "论文注意图号连续，不照搬原题图注。"),
    ("13", "附件总大小≤50MB",
     "100条特征+两组模型权重+视频极易超限。",
     "已控制。交付包26MB，关键帧降采样JPG。"),
    ("14", "label-100.xlsx残留出题人路径",
     "修复说明表里留有出题人本机桌面路径，说明该附件是二次加工且已按本地100片段筛选，与附件2的4850条不是同一集合，不能混用做训练集。",
     "正确。Q1自提取特征独立，不混用附件1的100条做训练。"),
]

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
headers = ["编号", "坑点", "具体内容", "本队应对"]
widths = [Cm(1.2), Cm(3.5), Cm(7.5), Cm(5.0)]
for i, h in enumerate(headers):
    set_cell(hdr[i], h, bold=True, size=10.5, bg="D9D9D9")
    hdr[i].width = widths[i]

for num, name, desc, action in pits:
    row = table.add_row().cells
    set_cell(row[0], num, size=10)
    set_cell(row[1], name, bold=True, size=10)
    set_cell(row[2], desc, size=10)
    set_cell(row[3], action, size=10)
    for i in range(4):
        row[i].width = widths[i]

doc.add_paragraph()

# ============ 第二部分：共同规律与执行建议 ============
add_heading("二、两题共同规律与执行建议", level=1)

add_heading("2.1 共同规律", level=2)
p = doc.add_paragraph()
p.add_run("两题都在'正文文本不完整或自相矛盾'上做文章，逼参赛队回到原始附件逐项核验；两题都埋了'看起来能用、其实用不了'的数据；两题都有题号/表号/图号缺陷，直接照抄会踩排版红线。能把这些矛盾写进论文的'数据说明'与'模型假设'，本身就是加分项。")

add_heading("2.2 E题执行建议", level=2)
for item in [
    "先统一输入接口（全程 text_bert）→ 跑通附件3/附件4完整推理管线 → 再做可解释性。",
    "若时间不足，优先保证'结果文件格式正确 + 基础指标可复现'，其次才是创新点。",
    "各留一节'数据核查与口径说明'，把实测矛盾（附件3字段缺失、掩码分布等）写清楚，并给出本队处理规则与依据。",
]:
    p = doc.add_paragraph(item, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.74)

doc.add_paragraph()

# ============ 第三部分：思路分析 ============
add_heading("三、E题思路分析（数据实况→三问方案）", level=1)

add_heading("3.1 数据实况（已逐项核查）", level=2)

data_table = doc.add_table(rows=1, cols=3)
data_table.style = "Table Grid"
hdr = data_table.rows[0].cells
for i, h in enumerate(["附件", "内容", "关键字段"]):
    set_cell(hdr[i], h, bold=True, bg="D9D9D9")

data_rows = [
    ("附件1", "37个video_id文件夹，100个mp4；label-100.xlsx共100行，label∈[-3,3]", "Positive/Neutral/Negative"),
    ("附件2", "aligned_50.pkl / unaligned_50.pkl；train 3395 / valid 728 / test 727 = 4850",
     "text(50×768)、audio(50×74)、vision(50×35)、text_bert(3×50)、classification_labels、regression_labels；类别分布 Positive 2370 / Negative 1380 / Neutral 1100"),
    ("附件3", "对齐/未对齐各30个pkl，每文件1条样本", "对齐版仅 text_bert/audio/vision；未对齐版仅 raw_text/audio/vision"),
    ("附件4", "对齐/未对齐各20个pkl + 20个mp4，字段完整（含id与text）", "可解释专项，配套视频可回看证据"),
]
for a, b, c in data_rows:
    row = data_table.add_row().cells
    set_cell(row[0], a, bold=True, size=10)
    set_cell(row[1], b, size=10)
    set_cell(row[2], c, size=10)

doc.add_paragraph()

add_heading("3.2 问题1：特征提取与时序对齐", level=2)
for item in [
    "从100条原始mp4出发，定义文本/语音/视觉三模态特征与统一对齐规则。",
    "统一到固定时间网格（如50段），每段做均值/最大池化。",
    "文本用BERT类编码器取token序列，同时保留(3×50)的token/attention/segment三路输入，以便对齐附件2接口。",
    "明确记录每条的原始有效时长、特征维度、对齐粒度与填充规则。",
    "交付物：100条特征文件 + 全量汇总表 + 至少1个典型样本的三模态时间对应图 + 工具版本与复现说明。",
]:
    doc.add_paragraph(item, style="List Bullet").paragraph_format.left_indent = Cm(0.74)

add_heading("3.3 问题2：局部缺失下的鲁棒预测", level=2)
for item in [
    "关键认知：这是时间轴上的连续段缺失，不是整模态缺失。",
    "模型：时间维掩码 + 模态门控融合；缺失段用插值或可学习token填补。",
    "训练期做'随机连续段掩码'增强，且掩码长度与位置分布要与附件3实测分布对齐。",
    "损失：分类交叉熵 + 回归MSE（或CCC）多任务加权；模型选择与阈值只在验证集上做。",
    "分析维度：缺失模态类型 × 缺失率 × 缺失位置。",
    "输出附件3预测CSV。",
]:
    doc.add_paragraph(item, style="List Bullet").paragraph_format.left_indent = Cm(0.74)

add_heading("3.4 问题3：可解释性预测", level=2)
for item in [
    "需同时给出极性、强度、模态作用程度、主要参考模态与局部关键证据定位。",
    "三路互证：跨模态注意力权重 + 时间维重要性（梯度/扰动/attention rollout）+ 反事实模态消融。",
    "避免只用注意力当解释。",
    "最硬的一环是证据回落：附件2不保存原始视频帧、音频波形或逐词时间戳，必须自行建立'特征位置→原始时间'的映射（附件4配套原始视频）。",
    "输出附件4预测+解释CSV。",
]:
    doc.add_paragraph(item, style="List Bullet").paragraph_format.left_indent = Cm(0.74)

doc.add_paragraph()

# ============ 第四部分：本队实际落地对照 ============
add_heading("四、本队实际落地与建议方案对照", level=1)

cmp_table = doc.add_table(rows=1, cols=3)
cmp_table.style = "Table Grid"
hdr = cmp_table.rows[0].cells
for i, h in enumerate(["维度", "网上建议方案", "本队实际实现"]):
    set_cell(hdr[i], h, bold=True, bg="D9D9D9")

cmp_rows = [
    ("输入接口", "全程 text_bert + 同一BERT编码器", "✅ 全程 text_bert attention_mask，附件3推理已修"),
    ("缺失增强", "连续段掩码，长度位置对齐附件3实测", "✅ audio+vision同时短缺失(长度1-4)，text不做中间缺失"),
    ("Q2模型", "时间维掩码+模态门控融合", "✅ ModalEncoder(Transformer,d=64)+门控融合+Huber+3·tanh"),
    ("Q2指标", "Acc/F1/MAE/Pearson", "✅ test Acc=0.663, F1=0.638, MAE=0.711, Pearson=0.642"),
    ("Q3解释", "注意力+梯度+遮挡三路互证", "✅ 注意力α/β + 遮挡法 + 积分梯度IG"),
    ("证据回落", "特征位置→原始时间映射", "✅ 附件4视频抽关键帧，IG定位时间步"),
    ("外部数据", "禁止外部情感数据集", "✅ 仅用附件2 CMU-MOSEI"),
    ("附件大小", "≤50MB", "✅ 交付包26MB"),
    ("创新点", "—", "AdaptiveSwish消融、知识蒸馏教师-学生(d=64→d=32)"),
]
for a, b, c in cmp_rows:
    row = cmp_table.add_row().cells
    set_cell(row[0], a, bold=True, size=10)
    set_cell(row[1], b, size=10)
    set_cell(row[2], c, size=10)

out = r"D:\Projects\math_modling\E题_坑点与思路分析.docx"
doc.save(out)
print(f"saved -> {out}")
